"""Lectura de hojas de vida: extracción de texto y lector por reglas (sin IA)."""
import pytest

from app.services import cv_parser_service
from app.services.cv_parser_service import CVReadError, extract_text, heuristic_parse, normalize
from tests.cv_samples import CV_LINES, make_docx, make_pdf

CV_TEXT = "\n".join(CV_LINES)


@pytest.mark.parametrize("ext,maker", [(".docx", make_docx), (".pdf", make_pdf)])
def test_extract_text_from_files(ext, maker):
    text = extract_text(maker(), ext)
    assert "Soluciones Andinas" in text
    assert "laura.gomez@correo.com" in text


def test_extract_text_rejects_empty_document():
    with pytest.raises(CVReadError):
        extract_text(make_pdf(["Hola"]), ".pdf")


def test_heuristic_parse_extracts_key_data():
    data = heuristic_parse(CV_TEXT)
    assert data["first_name"] == "Laura Marcela"
    assert data["last_name"] == "Gómez Ríos"
    assert data["email"] == "laura.gomez@correo.com"
    assert "301" in data["phone"]
    assert data["summary"].startswith("Ingeniera de sistemas")

    first, second = data["experiences"][:2]
    assert first["position"] == "Desarrolladora Backend"
    assert first["company_name"] == "Soluciones Andinas S.A.S."
    assert first["start_date"] == "2022-01" and first["is_current"] is True
    assert second["company_name"] == "Tecnología del Pacífico"
    assert (second["start_date"], second["end_date"]) == ("2020-03", "2021-12")

    edu = data["educations"][0]
    assert edu["degree"] == "Ingeniería de Sistemas"
    assert edu["institution"] == "Universidad del Valle"
    assert edu["level"] == "pregrado"

    assert {"Python", "FastAPI", "MySQL"} <= set(data["skills"])
    assert {"name": "ingles", "proficiency": "intermedio"} in data["languages"]


def test_normalize_maps_catalog_ids(db_session):
    with db_session() as db:
        data = normalize(db, heuristic_parse(CV_TEXT))
    assert data["city_id"] == 1  # Cali
    assert data["educations"][0]["education_level_id"] == 4  # Pregrado
    assert {"language_id": 2, "proficiency": "intermedio"} in data["languages"]
    assert data["experiences"][0]["start_date"] == "2022-01-01"


def test_ai_failure_falls_back_to_rules(db_session, monkeypatch):
    class BrokenAI:
        name = "falla"

        def complete_json(self, system, user):
            raise cv_parser_service.AIError("sin conexión")

    monkeypatch.setattr(cv_parser_service, "get_ai_provider", lambda: BrokenAI())
    with db_session() as db:
        data, source = cv_parser_service.parse_cv(db, CV_TEXT)
    assert source == "basico"
    assert data["email"] == "laura.gomez@correo.com"


def test_claude_provider_used_when_configured(db_session, monkeypatch):
    """Con AI_PROVIDER=claude se llama al SDK (simulado aquí) y se normaliza su JSON."""
    from types import SimpleNamespace

    from app.config import settings
    from app.services.ai.claude_provider import ClaudeProvider

    monkeypatch.setattr(settings, "AI_PROVIDER", "claude")
    monkeypatch.setattr(settings, "ANTHROPIC_API_KEY", "clave-de-prueba")
    reply = (
        '{"first_name": "Laura", "last_name": "Gómez", "email": "laura@correo.com", "phone": null,'
        ' "city": "Bogotá", "desired_position": "Backend", "summary": null, "skills": ["Python"],'
        ' "experiences": [], "languages": [],'
        ' "educations": [{"institution": "Univalle", "degree": "Ingeniería", "field_of_study": null,'
        ' "level": "Pregrado", "start_date": "2014-02", "end_date": "2019-12"}]}'
    )
    calls = []

    def fake_create(**kwargs):
        calls.append(kwargs)
        return SimpleNamespace(stop_reason="end_turn", content=[SimpleNamespace(type="text", text=reply)])

    monkeypatch.setattr(ClaudeProvider, "__init__", lambda self: setattr(self, "model", "m") or setattr(
        self, "client", SimpleNamespace(beta=SimpleNamespace(messages=SimpleNamespace(create=fake_create)))))

    with db_session() as db:
        data, source = cv_parser_service.parse_cv(db, CV_TEXT)
    assert source == "ia" and calls
    assert data["city_id"] == 2 and data["educations"][0]["education_level_id"] == 4
    assert data["educations"][0]["start_date"] == "2014-02-01"


# =====================================================================
# Hojas de vida con otra estructura (como las que exportan Canva / Word / PDF de dos niveles)
# Datos inventados: sirve para no volver a dañar lo que se corrigió.
# =====================================================================
STUDENT_CV = """Mariana Torres Ruiz — Curriculum Vitae
Página 1 de 1
MARIANA TORRES RUIZ
Calle 5 # 20-30, Cali, Colombia | Teléfono: 3001234567 | mariana.torres@correo.com
PERFIL PROFESIONAL
Estudiante de Ingeniería de Sistemas con enfoque en desarrollo de software.
Certificada en inglés B2.
EDUCACIÓN Y CERTIFICACIONES
Universidad del Valle
Cali, Colombia
Ingeniería de Sistemas (Quinto Semestre)
Grado esperado: 2029
PROYECTOS DE SOFTWARE
Tienda Online – Proyecto de clase
 Tecnologías: Django, JavaScript, SQL, Chart.js. Repositorio: github.com/mariana/tienda
 Desarrollo de módulos de inventario y caja.
HABILIDADES TÉCNICAS E IDIOMAS
Lenguajes de Programación: Python, Java, HTML, CSS.
Bases de Datos y Control: MySQL, Git, GitHub.
Infraestructura y Herramientas: Fundamentos de Linux, Redes (Cisco), Microsoft Office, Herramientas IA (Optimización de
Prompts).
Idiomas: Español (Nativo), Inglés (B2 Certificado - Fluidez técnica y profesional).
HABILIDADES BLANDAS
Liderazgo y Gestión: Capacidad para tomar la iniciativa en proyectos, estructurar tareas y organizar requerimientos técnicos.
Comunicación y Organización: Claridad en la transmisión de ideas, documentación estructurada y adaptabilidad al trabajo bajo
presión.
"""


def test_student_cv_with_combined_headers(db_session):
    from app.services.cv_parser_service import _clean_text

    with db_session() as db:
        data = normalize(db, heuristic_parse(_clean_text(STUDENT_CV)))
    assert (data["first_name"], data["last_name"]) == ("Mariana", "Torres Ruiz")
    assert data["phone"] == "300 123 4567" and data["city_id"] == 1

    skills = data["skills"]
    for expected in ("Python", "Java", "MySQL", "Git", "Linux", "Cisco", "Microsoft Office", "Django", "Chart.js",
                     "Liderazgo", "Comunicación", "Trabajo bajo presión", "Ingeniería de prompts"):
        assert expected in skills, expected
    # Nada de frases ni pedazos de líneas cortadas
    for bad in ("Estructurar tareas", "Prompts)", "Herramientas IA (Optimización de", "Presión"):
        assert bad not in skills, bad
    assert len({s.lower() for s in skills}) == len(skills)  # sin duplicados

    # Cada idioma con su propio nivel (antes ambos quedaban "nativo")
    assert {"language_id": 1, "proficiency": "nativo"} in data["languages"]
    assert {"language_id": 2, "proficiency": "intermedio"} in data["languages"]

    edu = data["educations"]
    assert len(edu) == 1
    assert edu[0]["degree"] == "Ingeniería de Sistemas" and edu[0]["institution"] == "Universidad del Valle"
    assert edu[0]["end_date"] is None  # en curso


def test_normalize_validates_and_dedupes_each_field(db_session):
    raw = {
        "first_name": "JUAN CARLOS", "last_name": "P3rez", "email": "no-es-correo", "phone": "+57 (310) 555-1234",
        "city": "Bogotá D.C.", "desired_position": "DESARROLLADOR DE SOFTWARE", "summary": "  Hola   mundo  ",
        "skills": ["python", "Python", "PYTHON", "excel", "js", "comunicacion", "de la empresa", "a", "Liderazgo:", "x" * 80],
        "experiences": [
            {"position": "DESARROLLADOR", "company_name": "ACME SAS", "start_date": "2020-01", "is_current": True},
            {"position": "Desarrollador", "company_name": "Acme Sas", "start_date": "2020-01", "is_current": True},
            {"position": "Analista", "company_name": "Analista", "start_date": "2030-01", "end_date": "2019-01"},
        ],
        "educations": [
            {"degree": "Ingeniería", "institution": "Univalle", "level": "Pregrado"},
            {"degree": "INGENIERÍA", "institution": "UNIVALLE", "level": "Pregrado"},
        ],
        "languages": [{"name": "Inglés", "proficiency": "basico"}, {"name": "english", "proficiency": "avanzado"}],
    }
    with db_session() as db:
        data = normalize(db, raw)
    assert data["first_name"] == "Juan Carlos" and data["last_name"] is None  # con números no es un nombre
    assert data["email"] is None and data["phone"] == "+57 310 555 1234"
    assert data["desired_position"] == "Desarrollador de Software" and data["summary"] == "Hola mundo"
    assert data["skills"] == ["Python", "Excel", "JavaScript", "Comunicación", "Liderazgo"]  # "Liderazgo:" se limpia
    assert len(data["experiences"]) == 2
    assert data["experiences"][0]["position"] == "Desarrollador" and data["experiences"][0]["company_name"] == "Acme S.A.S."
    assert data["experiences"][1]["company_name"] is None and data["experiences"][1]["start_date"] is None  # fecha futura
    assert len(data["educations"]) == 1
    assert data["languages"] == [{"language_id": 2, "proficiency": "avanzado"}]


def _make_odt(lines):
    import io
    import zipfile

    body = "".join(f"<text:p>{line}</text:p>" for line in lines)
    xml = ('<?xml version="1.0" encoding="UTF-8"?><office:document-content '
           'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
           'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"><office:body><office:text>'
           f"{body}</office:text></office:body></office:document-content>")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        z.writestr("content.xml", xml)
    return buf.getvalue()


def test_other_formats_odt_rtf_txt():
    odt = extract_text(_make_odt(CV_LINES), ".odt")
    assert "Soluciones Andinas" in odt and "laura.gomez@correo.com" in odt

    rtf = (r"{\rtf1\ansi{\fonttbl{\f0 Arial;}}\f0 LAURA G\'d3MEZ\par Desarrolladora Backend\par "
           r"Correo: laura.gomez@correo.com\par Habilidades: Python, SQL, Excel\par Ingl\'e9s intermedio\par}").encode("latin-1")
    text = extract_text(rtf, ".rtf")
    assert "LAURA GÓMEZ" in text and "Habilidades: Python, SQL, Excel" in text and "Arial" not in text

    txt = "\n".join(CV_LINES).encode("cp1252")
    assert "Ingeniería de Sistemas" in extract_text(txt, ".txt")


def test_docx_reads_tables_in_order():
    import io

    from docx import Document

    doc = Document()
    doc.add_paragraph("ANA PÉREZ")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "HABILIDADES"
    table.rows[0].cells[1].text = "Excel, Contabilidad"
    doc.add_paragraph("Fin del documento")
    buf = io.BytesIO()
    doc.save(buf)
    text = extract_text(buf.getvalue(), ".docx")
    assert text.index("ANA PÉREZ") < text.index("HABILIDADES") < text.index("Fin del documento")
